"""Extract plain text from supported file types."""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

from pypdf import PdfReader

logger = logging.getLogger(__name__)

# Image formats supported when OCR is enabled.
OCR_IMAGE_EXTENSIONS: tuple[str, ...] = (
    ".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp",
)


def extract_text(
    path: str | Path,
    *,
    ocr_enabled: bool = False,
    ocr_lang: str = "eng",
    ocr_min_chars: int = 50,
) -> str:
    """Return the text contents of `path` based on its extension.

    Args:
        path: Path to the file.
        ocr_enabled: When True, image files are OCR'd and scanned PDFs (pages
            with fewer than *ocr_min_chars* extracted characters) fall back to
            OCR automatically.  Requires ``pytesseract`` and (for PDFs)
            ``pdf2image`` / ``poppler`` to be installed.
        ocr_lang: Tesseract language code(s), e.g. ``"eng"`` or ``"eng+deu"``.
        ocr_min_chars: Minimum characters per page to consider text "present".
            Pages below this threshold trigger the OCR fallback.
    """

    p = Path(path)
    ext = p.suffix.lower()

    if ext in (".txt", ".md"):
        return _read_utf8(p)
    if ext == ".pdf":
        return _read_pdf(
            p,
            ocr_enabled=ocr_enabled,
            ocr_lang=ocr_lang,
            ocr_min_chars=ocr_min_chars,
        )
    if ext == ".docx":
        return _read_docx(p)
    if ext in (".html", ".htm"):
        return _read_html(p)
    if ext == ".csv":
        return _read_csv(p)
    if ext in OCR_IMAGE_EXTENSIONS:
        if not ocr_enabled:
            raise ValueError(
                f"File type '{ext}' requires OCR.  Set ocr.enabled: true in "
                "config.yaml and install pytesseract + Tesseract."
            )
        return _ocr_image(p, lang=ocr_lang)
    raise ValueError(f"Unsupported file type for extraction: {ext}")


def _read_utf8(path: Path) -> str:
    # `errors="replace"` keeps the pipeline running on files with stray bytes.
    return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(
    path: Path,
    *,
    ocr_enabled: bool = False,
    ocr_lang: str = "eng",
    ocr_min_chars: int = 50,
) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page_num, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if ocr_enabled and len(text.strip()) < ocr_min_chars:
            ocr_text = _ocr_pdf_page(path, page_num=page_num, lang=ocr_lang)
            if ocr_text:
                logger.debug(
                    "PDF page %d in %s used OCR fallback (%d chars extracted vs "
                    "%d threshold).",
                    page_num,
                    path.name,
                    len(text.strip()),
                    ocr_min_chars,
                )
                text = ocr_text
        pages.append(text)
    return "\n\n".join(pages)


def _ocr_pdf_page(path: Path, *, page_num: int, lang: str) -> str:
    """Render a single PDF page to an image and OCR it.

    Returns empty string if ``pdf2image`` or ``pytesseract`` are not installed.
    """
    try:
        from pdf2image import convert_from_path  # type: ignore[import-untyped]
    except ImportError:
        logger.warning(
            "pdf2image is not installed — OCR fallback unavailable for %s. "
            "Run: pip install pdf2image",
            path.name,
        )
        return ""
    try:
        import pytesseract  # type: ignore[import-untyped]
    except ImportError:
        logger.warning(
            "pytesseract is not installed — OCR fallback unavailable for %s. "
            "Run: pip install pytesseract  (and install Tesseract on your OS)",
            path.name,
        )
        return ""

    images = convert_from_path(str(path), first_page=page_num + 1, last_page=page_num + 1)
    if not images:
        return ""
    return pytesseract.image_to_string(images[0], lang=lang)


def _ocr_image(path: Path, *, lang: str) -> str:
    """OCR a standalone image file and return the extracted text."""
    try:
        from PIL import Image  # type: ignore[import-untyped]
        import pytesseract  # type: ignore[import-untyped]
    except ImportError as exc:
        raise ImportError(
            "OCR for image files requires Pillow and pytesseract: "
            "pip install Pillow pytesseract"
        ) from exc

    img = Image.open(str(path))
    return pytesseract.image_to_string(img, lang=lang)



def _read_docx(path: Path) -> str:
    """Extract text from a .docx file using python-docx."""

    from docx import Document  # type: ignore[import-untyped]

    doc = Document(str(path))
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)
    # Also extract text from tables.
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n\n".join(paragraphs)


def _read_html(path: Path) -> str:
    """Extract text from HTML using BeautifulSoup."""

    from bs4 import BeautifulSoup  # type: ignore[import-untyped]

    raw = path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")

    # Remove script and style elements.
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    text = soup.get_text(separator="\n", strip=True)
    # Collapse multiple blank lines.
    lines = [line for line in text.splitlines() if line.strip()]
    return "\n\n".join(lines)


def _read_csv(path: Path) -> str:
    """Convert a CSV file into readable text rows."""

    raw = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.reader(io.StringIO(raw))
    rows: list[str] = []
    header: list[str] | None = None
    for i, row in enumerate(reader):
        if i == 0:
            header = row
            rows.append(" | ".join(row))
        elif header:
            # Format as "Header1: val1, Header2: val2, ..."
            pairs = [f"{h}: {v}" for h, v in zip(header, row) if v.strip()]
            rows.append(", ".join(pairs))
        else:
            rows.append(" | ".join(row))
    return "\n".join(rows)
